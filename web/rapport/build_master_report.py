# -*- coding: utf-8 -*-
"""Rapport académique de synthèse — Strix + RAG + Metasploit + Console web."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(__file__)
FIG = os.path.join(HERE, "figures")                 # figures console
STRIXFIG = os.path.join(HERE, "strix", "figures")   # figures Strix
RAGFIG = os.path.join(HERE, "rag", "figures")       # figures RAG
SHOTS = os.path.join(HERE, "screenshots")           # captures interface
NAVY = RGBColor(0x0F, 0x24, 0x2E); BLUE = RGBColor(0x1E, 0x3A, 0x8A); GREY = RGBColor(0x47, 0x55, 0x69)

doc = Document()
normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.3
for lvl, sz, col in [("Heading 1", 18, BLUE), ("Heading 2", 14, NAVY), ("Heading 3", 12, GREY)]:
    st = doc.styles[lvl]; st.font.name = "Calibri"; st.font.size = Pt(sz); st.font.color.rgb = col; st.font.bold = True
sec = doc.sections[0]; sec.left_margin = sec.right_margin = Inches(1.0); sec.top_margin = sec.bottom_margin = Inches(1.0)

_fig = [0]        # numérotation continue des figures
_tbl = [0]        # numérotation continue des tableaux


def p(text="", size=11, bold=False, italic=False, align=None, color=None, after=6):
    para = doc.add_paragraph(); para.paragraph_format.space_after = Pt(after)
    if align: para.alignment = align
    r = para.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    return para


def h1(t): doc.add_page_break(); doc.add_heading(t, level=1)
def h2(t): doc.add_heading(t, level=2)
def h3(t): doc.add_heading(t, level=3)


def bullets(items):
    for it in items:
        para = doc.add_paragraph(style="List Bullet"); para.paragraph_format.space_after = Pt(3)
        if isinstance(it, tuple):
            para.add_run(it[0]).bold = True; para.add_run(" " + it[1])
        else:
            para.add_run(it)


def code(text):
    c = doc.add_paragraph(); c.paragraph_format.left_indent = Inches(0.3)
    c.paragraph_format.space_after = Pt(8)
    r = c.add_run(text); r.font.name = "Consolas"; r.font.size = Pt(9.5)


def figure(name, caption, base=FIG, width=6.2):
    _fig[0] += 1
    doc.add_picture(os.path.join(base, name), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(f"Figure {_fig[0]} — {caption}"); r.italic = True
    r.font.size = Pt(9.5); r.font.color.rgb = GREY; cap.paragraph_format.space_after = Pt(12)


def table(headers, rows, widths=None, caption=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htext in enumerate(headers):
        run = t.rows[0].cells[i].paragraphs[0].add_run(htext); run.bold = True; run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].paragraphs[0].add_run(str(val)).font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    if caption:
        _tbl[0] += 1
        cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(f"Tableau {_tbl[0]} — {caption}"); r.italic = True
        r.font.size = Pt(9.5); r.font.color.rgb = GREY
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def toc_field(title, kind="TOC"):
    doc.add_heading(title, level=1)
    para = doc.add_paragraph(); run = para.add_run()
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    b = OxmlElement("w:instrText"); b.set(qn("xml:space"), "preserve")
    b.text = 'TOC \\o "1-3" \\h \\z \\u' if kind == "TOC" else f'TOC \\h \\z \\c "{kind}"'
    c = OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"), "separate")
    d = OxmlElement("w:t"); d.text = "Clic droit → « Mettre à jour les champs »."
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for el in (a, b, c, d, e): run._r.append(el)


def part(title, subtitle):
    doc.add_page_break()
    p("", after=60)
    p(title, 26, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE, after=10)
    p(subtitle, 13, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY)


# ══════════════════════════════════════════════════════════════════════════════
# PAGE DE GARDE
# ══════════════════════════════════════════════════════════════════════════════
p("République Tunisienne", 11, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
p("Ministère de l'Enseignement Supérieur et de la Recherche Scientifique", 10.5,
  align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
p("Université / École d'Ingénieurs — [Nom de l'établissement]", 10.5,
  align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
p("RAPPORT DE PROJET DE FIN D'ÉTUDES", 14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY, after=4)
p("En vue de l'obtention du Diplôme National d'Ingénieur", 11, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
p("Spécialité : Sécurité des Systèmes Informatiques et des Réseaux", 11, italic=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, after=28)
p("Plateforme de tests d'intrusion autonome pilotée par IA agentique :", 16, bold=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE, after=2)
p("intégration du moteur Strix, d'un service RAG et de Metasploit,", 16, bold=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE, after=2)
p("et conception d'une console web d'opérations de sécurité (Strix Console)", 16, bold=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE, after=30)
p("Réalisé par : [Nom et Prénom de l'étudiant]", 12, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
p("Encadrant académique : [Nom]", 11, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
p("Encadrant professionnel : [Nom]", 11, align=WD_ALIGN_PARAGRAPH.CENTER, after=28)
p("Année Universitaire 2025 – 2026", 12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

# ── Dédicace ──────────────────────────────────────────────────────────────────
doc.add_page_break(); doc.add_heading("Dédicace", level=1)
p("À mes chers parents, pour leur amour inconditionnel, leurs sacrifices et leur confiance "
  "indéfectible qui ont été ma plus grande source de force tout au long de ce parcours.")
p("À ma famille et à mes proches, pour leur soutien constant et leurs encouragements.")
p("À mes enseignants et encadrants, pour la qualité de leur accompagnement et la rigueur qu'ils "
  "m'ont transmise.")
p("À mes amis et camarades, pour tous les moments partagés qui ont donné du sens à cette "
  "expérience.")

# ── Remerciements ─────────────────────────────────────────────────────────────
doc.add_page_break(); doc.add_heading("Remerciements", level=1)
p("Au terme de ce travail, je tiens à exprimer ma profonde gratitude à toutes les personnes qui "
  "ont contribué, de près ou de loin, à l'aboutissement de ce projet de fin d'études.")
p("Je remercie chaleureusement mon encadrant académique pour sa disponibilité, ses conseils "
  "avisés et son suivi rigoureux, ainsi que mon encadrant professionnel pour la confiance qu'il "
  "m'a accordée et pour le cadre technique stimulant qu'il a su instaurer.")
p("Mes remerciements s'adressent également aux membres du jury pour l'honneur qu'ils me font en "
  "acceptant d'évaluer ce travail, ainsi qu'à l'ensemble du corps enseignant pour la formation de "
  "qualité dont j'ai bénéficié.")

# ── Résumé / Abstract ─────────────────────────────────────────────────────────
doc.add_page_break(); doc.add_heading("Résumé", level=1)
p("Ce projet de fin d'études porte sur la conception et la réalisation d'une plateforme complète "
  "de tests d'intrusion autonomes, pilotée par un système d'intelligence artificielle agentique "
  "fondé sur des grands modèles de langage (LLM). Le travail articule quatre briques "
  "complémentaires en un système cohérent. (1) Le moteur agentique open-source Strix orchestre "
  "des agents qui exécutent des outils offensifs réels dans un bac à sable Docker/Kali. (2) Un "
  "service de connaissances RAG (Retrieval-Augmented Generation) auto-hébergé — FastAPI, base "
  "vectorielle Qdrant, embeddings et reranking neuronal — ancre le raisonnement des agents dans "
  "un corpus offensif vérifié afin de réduire les hallucinations. (3) Le framework d'exploitation "
  "Metasploit apporte un arsenal éprouvé ; un outil unique et déterministe, exploit_research, "
  "collapse le protocole RAG + Metasploit en un seul appel et produit une recommandation "
  "« module d'abord », déplaçant la décision d'exploitation du modèle vers du code vérifiable. "
  "(4) Une console web professionnelle, Strix Console (Next.js + FastAPI), transforme cet "
  "outillage en un véritable service d'opérations de sécurité : authentification forte (cookies "
  "httpOnly, CSRF, MFA TOTP), contrôle d'accès par rôles (RBAC), pilotage et suivi temps réel des "
  "campagnes (WebSocket), visualisation de la sortie du moteur, post-exploitation, triage des "
  "vulnérabilités, comparaison de scans, planification récurrente, export de rapports et journal "
  "d'audit inaltérable. La plateforme a été validée par 39 tests automatisés et vérifiée de bout "
  "en bout.")
p("Mots-clés : ", bold=True, after=0)
p("tests d'intrusion, IA agentique, LLM, Strix, RAG, Qdrant, Metasploit, RBAC, MFA, FastAPI, "
  "Next.js, sécurité applicative.", italic=True)
doc.add_heading("Abstract", level=1)
p("This graduation project designs and builds a complete platform for autonomous penetration "
  "testing, driven by an agentic AI system based on Large Language Models (LLMs). It integrates "
  "four complementary building blocks into a coherent system. (1) The open-source Strix agentic "
  "engine orchestrates agents that run real offensive tools inside a Docker/Kali sandbox. (2) A "
  "self-hosted Retrieval-Augmented Generation (RAG) knowledge service — FastAPI, Qdrant vector "
  "store, neural embeddings and reranking — grounds the agents' reasoning in a verified offensive "
  "corpus to reduce hallucinations. (3) The Metasploit framework provides a battle-tested "
  "arsenal; a single deterministic tool, exploit_research, collapses the RAG + Metasploit "
  "protocol into one call and returns a module-first recommendation, moving exploitation "
  "decisions from the model to verifiable code. (4) A professional web console, Strix Console "
  "(Next.js + FastAPI), turns this tooling into a real security-operations service: strong "
  "authentication, RBAC, real-time orchestration (WebSocket), engine-output visualisation, "
  "post-exploitation, findings triage, scan diffing, recurring scheduling, report export and a "
  "tamper-evident audit trail. The platform is covered by 39 automated tests and verified "
  "end-to-end.")
p("Keywords: ", bold=True, after=0)
p("penetration testing, agentic AI, LLM, Strix, RAG, Qdrant, Metasploit, RBAC, MFA, FastAPI, "
  "Next.js.", italic=True)

# ── Tables ────────────────────────────────────────────────────────────────────
doc.add_page_break(); toc_field("Table des matières", "TOC")
doc.add_page_break(); toc_field("Liste des figures", "Figure")
doc.add_page_break(); toc_field("Liste des tableaux", "Tableau")

# ── Abréviations ──────────────────────────────────────────────────────────────
doc.add_page_break(); doc.add_heading("Liste des abréviations", level=1)
table(["Abréviation", "Signification"], [
    ["LLM", "Large Language Model (grand modèle de langage)"],
    ["RAG", "Retrieval-Augmented Generation (génération augmentée de récupération)"],
    ["MSF", "Metasploit Framework"],
    ["RBAC", "Role-Based Access Control (contrôle d'accès par rôles)"],
    ["MFA / TOTP", "Multi-Factor Authentication / Time-based One-Time Password"],
    ["CSRF", "Cross-Site Request Forgery"],
    ["API", "Application Programming Interface"],
    ["WS", "WebSocket"],
    ["CVE", "Common Vulnerabilities and Exposures"],
    ["OCR", "Optical Character Recognition (reconnaissance optique de caractères)"],
    ["ORM", "Object-Relational Mapping"],
    ["HIL", "Human-In-the-Loop (validation humaine)"],
    ["CI/CD", "Continuous Integration / Continuous Deployment"],
    ["PoC", "Proof of Concept (preuve de concept)"],
], widths=[1.5, 4.8])

# ══════════════════════════════════════════════════════════════════════════════
# INTRODUCTION GÉNÉRALE
# ══════════════════════════════════════════════════════════════════════════════
h1("Introduction générale")
p("La transformation numérique des organisations s'accompagne d'une surface d'attaque en "
  "expansion permanente. Les tests d'intrusion (penetration testing) constituent l'un des moyens "
  "les plus efficaces pour évaluer, de façon offensive et contrôlée, la robustesse d'un système "
  "d'information. Toutefois, réalisés manuellement, ils demeurent coûteux en temps, difficilement "
  "reproductibles et fortement dépendants de l'expertise de l'opérateur.")
p("L'émergence des grands modèles de langage (LLM) et des architectures dites « agentiques » "
  "ouvre la voie à une automatisation intelligente de ce processus : des agents autonomes "
  "planifient, exécutent et enchaînent des actions offensives en s'appuyant sur des outils réels. "
  "Le moteur open-source Strix illustre cette approche. Néanmoins, deux limites subsistent : "
  "d'une part, la phase d'exploitation reste fragile lorsqu'elle est entièrement confiée au "
  "raisonnement du modèle, sujet aux hallucinations ; d'autre part, l'usage en équipe manque "
  "d'une interface d'opérations sécurisée offrant gouvernance, traçabilité et pilotage.")
p("Ce projet répond à ces deux limites par une approche hybride et intégrée. Il ancre le "
  "raisonnement des agents dans un service de connaissances RAG spécialisé pour l'exploitation, "
  "privilégie un exécuteur éprouvé (Metasploit) au moyen d'un outil déterministe unique, puis "
  "conçoit et réalise une console web professionnelle, Strix Console, dédiée à l'orchestration et "
  "au suivi sécurisés des campagnes. Le résultat forme une chaîne cohérente allant de "
  "l'interface d'opérations jusqu'à l'exécution offensive dans le bac à sable.")
p("Le présent rapport est organisé en six parties. La première pose le cadre du projet et l'état "
  "de l'art. La deuxième analyse le moteur agentique Strix. La troisième détaille le service de "
  "connaissances Pentest RAG. La quatrième présente l'intégration Metasploit + RAG qui fiabilise "
  "l'exploitation. La cinquième expose la conception, la réalisation et l'interface de la console "
  "web. La sixième présente les tests et l'évaluation. Une conclusion générale synthétise les "
  "apports et ouvre des perspectives.")

# ══════════════════════════════════════════════════════════════════════════════
# PARTIE I
# ══════════════════════════════════════════════════════════════════════════════
part("Partie I", "Contexte, problématique et état de l'art")

h1("Chapitre 1 — Cadre général du projet")
h2("1.1 Introduction")
p("Ce chapitre situe le projet dans son contexte, formule la problématique, présente la solution "
  "proposée et décrit la méthodologie de gestion adoptée pour mener le développement.")
h2("1.2 Contexte et problématique")
p("Le pentest manuel souffre de plusieurs limites structurelles : durée d'exécution élevée, "
  "dépendance à l'expertise humaine, faible reproductibilité et difficulté de traçabilité. Les "
  "plateformes d'automatisation par IA agentique, comme Strix, apportent une réponse prometteuse "
  "mais présentent, en l'état, deux insuffisances :")
bullets([
    ("Exploitation fragile — ", "en confiant l'intégralité de la décision d'exploitation au LLM, "
     "l'agent peut échouer à sélectionner ou paramétrer correctement un exploit, faute d'accès "
     "structuré à une base de connaissances et à un arsenal éprouvé."),
    ("Absence d'interface d'opérations — ", "l'outil est piloté en ligne de commande, sans "
     "gestion des utilisateurs, sans contrôle d'accès, sans suivi temps réel ni gouvernance, ce "
     "qui limite l'usage collaboratif et la valeur opérationnelle en entreprise."),
])
p("La problématique retenue peut ainsi se formuler : comment fiabiliser la phase d'exploitation "
  "d'un agent de pentest fondé sur un LLM, et comment en faire un service d'opérations de "
  "sécurité sûr, gouverné et exploitable en équipe ?")
h2("1.3 Solution proposée et contributions")
p("La solution combine quatre briques en un système intégré. Les contributions propres de ce "
  "projet portent sur l'intégration (Metasploit + RAG) et sur la console web ; le moteur Strix et "
  "le service RAG sont mobilisés et adaptés au service de cet objectif.")
table(["Brique", "Nature", "Apport dans le projet"], [
    ["Strix", "Moteur agentique open-source (mobilisé)", "Orchestration multi-agent et exécution des outils offensifs"],
    ["Pentest RAG", "Service de connaissances (conçu/adapté)", "Ancrage du raisonnement dans un corpus offensif vérifié"],
    ["Metasploit", "Framework d'exploitation (intégré)", "Exécuteur éprouvé, vérification non destructive (check)"],
    ["exploit_research", "Outil déterministe (contribution)", "Recommandation « module d'abord », décision codée en Python"],
    ["Strix Console", "Console web (contribution)", "Sécurité, RBAC, temps réel, gouvernance et reporting"],
], widths=[1.5, 2.3, 2.5], caption="Briques du système et contributions du projet.")
h2("1.4 Objectifs")
bullets([
    "Fiabiliser la phase d'exploitation en l'ancrant dans des connaissances vérifiées et un "
    "arsenal éprouvé.",
    "Rendre la décision d'exploitation déterministe et testable plutôt que laissée au seul modèle.",
    "Offrir une plateforme d'opérations sécurisée, gouvernée et collaborative.",
    "Assurer le pilotage et le suivi temps réel des campagnes, du lancement au reporting.",
    "Garantir la qualité par des tests automatisés et une vérification de bout en bout.",
])
h2("1.5 Méthodologie de gestion de projet")
p("Le développement a suivi une démarche itérative et incrémentale, organisée en phases "
  "successives livrant chacune un incrément fonctionnel testé et vérifié. Cette approche, proche "
  "de Scrum/Kanban, a permis de sécuriser les fondations avant d'empiler les fonctionnalités, "
  "tout en conservant une application exécutable à chaque étape.")
bullets([
    ("Intégration Metasploit + RAG — ", "outil exploit_research, playbooks et tests."),
    ("Phase 0 (Fondations) — ", "couche de données asynchrone, migrations, files et harnais de tests."),
    ("Phase 1 (Authentification) — ", "cookies httpOnly, CSRF, MFA TOTP, RBAC, sessions, audit."),
    ("Phase 2 (Temps réel) — ", "WebSocket, réconciliation des exécutions orphelines."),
    ("Phase 3 (Sortie moteur) — ", "arbre d'agents, console live, télémétrie (jetons, coût)."),
    ("Phase 4 (Post-exploitation) — ", "5ᵉ phase du pipeline et événements typés."),
    ("Phase 5 (Domaine) — ", "triage, comparaison de scans, alertes, planification récurrente."),
    ("Phase 6 (Qualité) — ", "palette de commandes, observabilité, durcissement."),
])
figure("gantt.png", "Planning du projet (méthodologie itérative par phases).")
h2("1.6 Conclusion")
p("Ce chapitre a défini le cadre, la problématique et la trajectoire du projet. Le chapitre "
  "suivant établit l'état de l'art sur lequel s'appuient les choix de conception.")

h1("Chapitre 2 — État de l'art")
h2("2.1 Introduction")
p("Ce chapitre présente les fondements des tests d'intrusion, les approches d'automatisation par "
  "IA agentique et les briques techniques mobilisées (Strix, RAG, Metasploit), afin de "
  "positionner la contribution du projet.")
h2("2.2 Tests d'intrusion et Red Teaming")
p("Un test d'intrusion reproduit, de manière autorisée et méthodique, le comportement d'un "
  "attaquant afin d'identifier et de démontrer des vulnérabilités exploitables. Le cycle de vie "
  "standard comprend la reconnaissance, l'évaluation des vulnérabilités, l'exploitation, la "
  "post-exploitation puis la rédaction du rapport. Le Red Teaming étend cette démarche à une "
  "simulation d'adversaire plus réaliste et furtive.")
h2("2.3 Automatisation par IA agentique et LLM")
p("Un système « agentique » orchestre un ou plusieurs agents fondés sur un LLM capables de "
  "raisonner, de décomposer une tâche, d'appeler des outils et d'itérer à partir des retours "
  "d'exécution. Appliqué au pentest, ce paradigme automatise la planification et l'enchaînement "
  "des actions offensives. Des travaux comme PentestGPT ou AutoGPT ont démontré la faisabilité de "
  "cette approche, tout en révélant ses limites : hallucinations, décisions d'exploitation peu "
  "fiables et besoin d'un ancrage dans des connaissances vérifiées.")
h2("2.4 Le paradigme RAG")
p("La génération augmentée de récupération (RAG) répond directement au problème de "
  "l'hallucination : plutôt que de s'appuyer sur la seule mémoire interne du modèle, le système "
  "récupère, dans une base de connaissances de confiance, les fragments pertinents pour la "
  "requête, puis les fournit au modèle comme contexte d'ancrage. Dans un contexte offensif, ce "
  "mécanisme évite les CVE inexistantes et les payloads erronés, en fondant chaque réponse sur "
  "des documents réels (writeups, exploits, données CVE).")
h2("2.5 Briques mobilisées")
bullets([
    ("Strix — ", "moteur multi-agent open-source de pentest ; un CLI côté hôte orchestre des "
     "agents qui exécutent des outils dans un bac à sable Docker (Kali)."),
    ("Pentest RAG — ", "service auto-hébergé (FastAPI + Qdrant) qui indexe des connaissances "
     "d'exploitation et fournit, sur requête, un guidage structuré ancré dans des documents réels."),
    ("Metasploit Framework — ", "arsenal d'exploitation de référence ; ses modules apportent "
     "vérification non destructive (check), gestion des cas limites et sessions structurées."),
])
p("La combinaison de ces briques permet d'ancrer le raisonnement de l'agent (RAG) tout en "
  "privilégiant un exécuteur éprouvé (Metasploit) lorsqu'un module existe, sous le pilotage d'une "
  "console d'opérations.")
h2("2.6 Positionnement du projet")
p("Par rapport à l'existant, la contribution est double : (i) rendre l'exploitation déterministe "
  "en collapsant le protocole RAG + Metasploit dans un outil unique dont la décision est codée en "
  "Python ; (ii) doter la plateforme d'une console d'opérations sécurisée, absente des outils "
  "comparables, apportant gouvernance, temps réel et traçabilité.")
h2("2.7 Conclusion")
p("L'état de l'art confirme la pertinence d'une approche hybride « IA agentique + connaissances "
  "vérifiées + arsenal éprouvé » et l'intérêt d'une interface d'opérations. Les parties suivantes "
  "détaillent chacune des briques, puis leur intégration.")

# ══════════════════════════════════════════════════════════════════════════════
# PARTIE II — STRIX
# ══════════════════════════════════════════════════════════════════════════════
part("Partie II", "Le moteur agentique Strix")

h1("Chapitre 3 — Architecture et fonctionnement de Strix")
h2("3.1 Introduction")
p("Strix est le socle offensif de la plateforme : c'est lui qui, in fine, exécute les actions "
  "d'attaque. Ce chapitre en présente l'architecture, le système multi-agent, la boîte à outils, "
  "la base de connaissances et le flux d'exécution.")
h2("3.2 Architecture technique")
p("L'architecture distingue l'orchestration (côté hôte) de l'exécution des outils (dans un bac à "
  "sable isolé). Un CLI Python orchestre les agents et présente une interface en terminal ; un "
  "bac à sable Docker fondé sur une image Kali exécute les outils via un serveur d'outils ; un "
  "fournisseur LLM est interrogé à travers la bibliothèque LiteLLM.")
figure("architecture.png", "Architecture technique du moteur Strix (hôte, bac à sable, LLM, sorties).",
       base=STRIXFIG)
h2("3.3 Système multi-agent (graphe d'agents)")
p("Plutôt qu'un agent unique, un agent racine décompose la cible en tâches indépendantes et "
  "engendre dynamiquement des sous-agents spécialisés (reconnaissance, injection, contrôle "
  "d'accès, SSRF/XXE, reporting…). Ces agents s'exécutent en parallèle, collaborent et partagent "
  "leurs découvertes ; des sous-agents de validation confirment l'exploitabilité par preuve de "
  "concept.")
figure("agents.png", "Graphe d'agents : orchestration multi-agent de Strix.", base=STRIXFIG)
h2("3.4 Boîte à outils des agents")
p("Les agents disposent d'un ensemble d'outils couvrant l'ensemble du cycle offensif, chacun "
  "exposé au LLM avec un schéma d'appel.")
figure("tools.png", "Catégories d'outils des agents Strix.", base=STRIXFIG, width=6.2)
table(["Outil", "Rôle"], [
    ["proxy (Caido)", "Proxy HTTP complet : interception et manipulation requêtes/réponses"],
    ["browser", "Automatisation d'un navigateur multi-onglets (XSS, CSRF, flux d'auth)"],
    ["terminal", "Shells interactifs pour l'exécution de commandes"],
    ["python", "Runtime Python pour développer et valider des exploits"],
    ["agents_graph", "Création et coordination des sous-agents"],
    ["reporting / notes / todo", "Documentation des vulnérabilités, gestion des découvertes et des tâches"],
    ["web_search", "Recherche d'informations (CVE, techniques, PoC)"],
    ["exploit_research", "(Contribution) RAG + Metasploit, recommandation « module d'abord »"],
    ["finish", "Clôture d'un agent ou de la campagne"],
], widths=[2.0, 4.3], caption="Principaux outils de la boîte à outils Strix.")
h2("3.5 Base de connaissances (skills)")
p("Les « skills » sont des paquets de connaissances spécialisées injectés dynamiquement dans le "
  "prompt d'un agent selon sa tâche (jusqu'à cinq par agent). Ils couvrent notamment 17 classes "
  "de vulnérabilités, des frameworks, des technologies, des protocoles, le cloud, la "
  "reconnaissance et les modes de scan.")
figure("vulns.png", "Classes de vulnérabilités couvertes par Strix (regroupées par famille).",
       base=STRIXFIG, width=6.0)
h2("3.6 Flux d'exécution et livrables")
p("Une campagne enchaîne reconnaissance, évaluation des vulnérabilités, exploitation, validation "
  "par PoC et reporting. Trois modes ajustent la profondeur (quick, standard, deep). Les "
  "résultats sont persistés dans un répertoire de campagne (strix_runs/) : vulnerabilities.csv, "
  "fiches par vulnérabilité, run.json (télémétrie LLM), journal d'exécution et rapport de "
  "synthèse.")
figure("workflow.png", "Flux d'exécution d'une campagne Strix.", base=STRIXFIG, width=6.5)
h2("3.7 Conclusion")
p("Strix fournit un moteur offensif puissant et extensible. Sa limite — la fragilité de "
  "l'exploitation confiée au modèle — motive l'ancrage par RAG et l'intégration de Metasploit, "
  "objets des parties suivantes.")

# ══════════════════════════════════════════════════════════════════════════════
# PARTIE III — RAG
# ══════════════════════════════════════════════════════════════════════════════
part("Partie III", "La base de connaissances Pentest RAG")

h1("Chapitre 4 — Conception du service Pentest RAG")
h2("4.1 Introduction")
p("Le service Pentest RAG ancre le raisonnement des agents dans un corpus offensif vérifié. Ce "
  "chapitre présente son architecture conteneurisée, son pipeline en trois phases et son API.")
h2("4.2 Architecture technique")
p("L'architecture est orchestrée par Docker Compose et réunit deux services sur un réseau privé : "
  "l'API applicative (FastAPI) et la base vectorielle Qdrant. Le backend LLM (Ollama en local ou "
  "OpenRouter en cloud) est sollicité en externe, et la persistance est assurée par des volumes.")
figure("architecture.png", "Architecture technique du service Pentest RAG.", base=RAGFIG)
table(["Composant", "Technologie"], [
    ["API applicative", "FastAPI + uvicorn (port 8000)"],
    ["Base vectorielle", "Qdrant (ports 6333/6334, distance COSINE)"],
    ["Embeddings", "nomic-ai/nomic-embed-text-v1"],
    ["Reranker", "cross-encoder/ms-marco-MiniLM-L-6-v2"],
    ["Backend LLM", "Ollama (local) / OpenRouter (cloud)"],
    ["Orchestration", "Docker Compose (réseau pentest_net)"],
], widths=[2.1, 4.2], caption="Composants du service RAG.")
h2("4.3 Pipeline en trois phases")
p("Le cœur fonctionnel repose sur trois phases adossées à Qdrant : l'ingestion construit la base "
  "vectorielle ; la récupération retrouve les fragments pertinents ; la génération produit une "
  "réponse structurée en JSON.")
figure("pipeline.png", "Pipeline du RAG en trois phases (ingestion, récupération, génération).",
       base=RAGFIG, width=6.5)
h3("4.3.1 Ingestion")
p("L'ingestion transforme des données brutes hétérogènes (CVE, writeups, exploits, images) en "
  "vecteurs indexés. Le pipeline nettoie le texte, extrait sélectivement les champs utiles du "
  "JSON CVE, applique l'OCR (Tesseract) aux images, filtre le bruit binaire et les fragments "
  "faibles, et déduplique à deux niveaux (fichier et chunk) par empreinte SHA-256. Les fragments "
  "(taille 500, chevauchement 200) sont vectorisés et insérés par lots robustes (réessais).")
h3("4.3.2 Récupération et reclassement")
p("La récupération combine expansion de requête, recherche sémantique large (top-100), "
  "reclassement par Cross-Encoder et scoring composite pondérant la pertinence sémantique, les "
  "correspondances lexicales, les signaux techniques (payload, CVE, curl…) et la correspondance "
  "au contexte cible (service, version, OS, findings). Les cinq meilleurs fragments forment le "
  "contexte.")
figure("retrieval.png", "Chaîne de récupération et de reclassement du RAG.", base=RAGFIG, width=6.2)
h3("4.3.3 Génération structurée")
p("Le modèle est contraint, par un prompt strict, à ne produire qu'un objet JSON conforme à un "
  "schéma imposé (StructuredRAGResponse), directement consommable par un agent.")
table(["Champ", "Description"], [
    ["vulnerability_id", "Identifiant précis (CVE-XXXX-XXXX ou nom standard)"],
    ["reasoning", "Explication technique de la faille et de son applicabilité"],
    ["actionable_commands", "Commandes ordonnées de vérification et de déclenchement"],
    ["payloads", "Chaînes de payload brutes (curl, versions encodées, reverse shells)"],
    ["exploit_code", "Script complet ou bloc de payload multi-lignes"],
    ["post_exploitation", "Persistance et élévation de privilèges"],
    ["confidence_score", "Score de fiabilité (0.0–1.0)"],
], widths=[1.9, 4.4], caption="Schéma de la réponse structurée du RAG.")
h2("4.4 API REST")
p("L'API expose trois endpoints : /retrieve (récupération rapide du contexte, sans LLM), /query "
  "(pipeline complet) et /ingest (indexation du corpus). Un cache en mémoire accélère les "
  "requêtes répétées.")
h2("4.5 Conclusion")
p("Le service RAG fournit un guidage d'exploitation fiable et structuré. Son intégration au "
  "moteur Strix, conjuguée à Metasploit, fait l'objet de la partie suivante.")

# ══════════════════════════════════════════════════════════════════════════════
# PARTIE IV — INTÉGRATION
# ══════════════════════════════════════════════════════════════════════════════
part("Partie IV", "Intégration Metasploit + RAG : fiabiliser l'exploitation")

h1("Chapitre 5 — L'outil déterministe exploit_research")
h2("5.1 Introduction")
p("Ce chapitre présente la contribution centrale du volet moteur : la fusion du protocole "
  "RAG + Metasploit en un unique outil déterministe, exploit_research, et le moteur "
  "d'exploitation à trois niveaux qui en découle.")
h2("5.2 Principe : décision « module d'abord »")
p("L'intégration transforme un protocole multi-étapes en un seul appel d'outil déterministe. Dès "
  "que la reconnaissance fournit un service, une version ou une CVE, l'agent invoque "
  "exploit_research, qui : (1) interroge le RAG (/retrieve puis, après un portail de pertinence, "
  "/query) en contournant le proxy et avec un délai adapté (120 s) ; (2) recherche un module via "
  "msfconsole ; (3) détecte l'adresse LHOST du bac à sable ; (4) produit une recommandation "
  "déterministe privilégiant le module Metasploit lorsqu'il existe, et retombant sur le payload "
  "RAG sinon. La logique de décision est ainsi codée et testée, et non laissée au modèle.")
figure("exploit_handoff.png", "Handoff d'exploitation : RAG puis Metasploit, avec priorité au module.")
h2("5.3 Moteur d'exploitation à trois niveaux")
p("La recommandation est exécutée par un moteur à repli hiérarchique, qui choisit la voie la plus "
  "fiable disponible selon trois niveaux de priorité décroissante.")
figure("tiers.png", "Moteur d'exploitation à trois niveaux avec barrière de validation.",
       base=RAGFIG, width=5.9)
table(["Niveau", "Condition", "Action"], [
    ["Tier 1 — MSF natif", "Module présent dans Metasploit", "Validation puis exécution via MSF RPC"],
    ["Tier 2 — SearchSploit", "Module absent mais trouvé via SearchSploit",
     "Import dans MSF, rechargement puis exécution"],
    ["Tier 3 — RAG direct", "Ni MSF ni SearchSploit",
     "Exécution des payloads HTTP (curl), du script, puis des commandes"],
], widths=[1.7, 2.3, 2.5], caption="Les trois niveaux d'exécution de l'exploitation.")
h2("5.4 Barrière de validation humaine (HIL)")
p("Quel que soit le niveau, aucun exploit n'est exécuté sans franchir une barrière de validation. "
  "Chaque plan démarre au statut « en attente d'approbation » et doit passer explicitement au "
  "statut « approuvé » avant toute exécution ; à défaut, il est bloqué. Ce garde-fou "
  "human-in-the-loop garantit qu'aucune action offensive n'est déclenchée de façon incontrôlée, "
  "et constitue un élément clé de la sécurité et de l'éthique du dispositif.")
h2("5.5 Protocole d'intégration côté agent")
p("Côté agent, le contrat d'appel est explicite : après reconnaissance, l'agent appelle "
  "d'abord /retrieve, applique un portail de pertinence structurel (contexte de plus de 100 "
  "caractères et d'au moins 3 lignes non vides), puis n'appelle /query que si le portail est "
  "franchi. Le champ confidence_score guide la décision : exploiter (> 0,7), vérifier d'abord "
  "(0,4–0,7) ou approfondir la reconnaissance (< 0,4).")
figure("strix_integration.png", "Protocole d'intégration RAG côté agent Strix.", base=RAGFIG, width=6.5)
code('# Construction de la commande du moteur (extrait)\n'
     'cmd = ["strix", "--target", target, "--non-interactive", "--scan-mode", mode]\n'
     'if instruction:\n'
     '    cmd += ["--instruction", instruction]')
h2("5.6 Conclusion")
p("L'intégration déplace la décision d'exploitation du modèle vers du code vérifiable et "
  "privilégie un exécuteur éprouvé, tout en conservant un repli fondé sur le RAG. La partie "
  "suivante présente la console qui pilote et gouverne l'ensemble.")

# ══════════════════════════════════════════════════════════════════════════════
# PARTIE V — CONSOLE
# ══════════════════════════════════════════════════════════════════════════════
part("Partie V", "La console web d'opérations — Strix Console")

h1("Chapitre 6 — Conception et architecture de la console")
h2("6.1 Introduction")
p("Ce chapitre décrit l'architecture globale de la console, les besoins fonctionnels et non "
  "fonctionnels, les diagrammes UML, le modèle de données et la conception de la sécurité.")
h2("6.2 Architecture globale")
p("La plateforme adopte une architecture trois tiers. Le client (Next.js) communique en même "
  "origine avec un serveur Next.js qui relaie les appels /api/* vers un backend FastAPI "
  "asynchrone. Ce dernier expose les services métier, persiste les données dans PostgreSQL "
  "(SQLite en développement) et pilote le moteur Strix, lequel s'appuie sur Metasploit et le "
  "service RAG. Redis fournit files et cache.")
figure("architecture.png", "Architecture globale de la plateforme Strix Console.")
h2("6.3 Besoins")
h3("6.3.1 Acteurs")
p("Quatre rôles système sont définis, du moins au plus privilégié : Observateur (lecture seule), "
  "Analyste (création et exécution de tests, triage, export), Manager (supervision de toutes les "
  "campagnes, audit, lecture des utilisateurs) et Administrateur (accès complet).")
figure("usecase.png", "Diagramme de cas d'utilisation de la plateforme.")
h3("6.3.2 Besoins fonctionnels")
table(["Réf.", "Besoin fonctionnel"], [
    ["BF1", "S'authentifier avec cookies sécurisés et second facteur (MFA TOTP)"],
    ["BF2", "Gérer les utilisateurs, les rôles et les permissions (RBAC)"],
    ["BF3", "Créer, lancer, annuler et supprimer un test d'intrusion"],
    ["BF4", "Suivre en temps réel le pipeline (état, phases, journal)"],
    ["BF5", "Visualiser la sortie du moteur (arbre d'agents, télémétrie, console)"],
    ["BF6", "Trier les vulnérabilités (statut, gravité, assignation)"],
    ["BF7", "Comparer deux scans d'une même cible (nouveaux / corrigés)"],
    ["BF8", "Planifier des tests récurrents (cron)"],
    ["BF9", "Exporter des rapports (PDF, CSV, JSON)"],
    ["BF10", "Consulter l'assistant IA et le journal d'audit"],
], widths=[0.8, 5.5], caption="Besoins fonctionnels principaux.")
h3("6.3.3 Besoins non fonctionnels")
table(["Réf.", "Besoin non fonctionnel"], [
    ["BNF1", "Sécurité : jetons httpOnly, protection CSRF, MFA, en-têtes de sécurité"],
    ["BNF2", "Performance : backend asynchrone, temps réel via WebSocket"],
    ["BNF3", "Traçabilité : journal d'audit inaltérable (chaîné par hachage)"],
    ["BNF4", "Fiabilité : réconciliation des exécutions, tests automatisés"],
    ["BNF5", "Maintenabilité : migrations versionnées, séparation claire des couches"],
    ["BNF6", "Observabilité : identifiants de requête, sondes /readyz et /livez"],
], widths=[0.8, 5.5], caption="Besoins non fonctionnels.")
h2("6.4 Diagramme de séquence — exécution d'un test")
p("Le scénario nominal d'exécution d'un test illustre l'enchaînement entre l'utilisateur, le "
  "frontend, l'API, le moteur et la base de données, ainsi que la diffusion des événements temps "
  "réel via WebSocket.")
figure("seq_run.png", "Diagramme de séquence : exécution d'un test.")
h2("6.5 Modèle de données")
p("Le modèle relationnel s'organise autour des entités suivantes, gérées via l'ORM SQLAlchemy "
  "2.0 et versionnées par Alembic.")
table(["Table", "Rôle"], [
    ["users / roles / role_permissions", "Comptes, rôles et permissions (RBAC)"],
    ["user_sessions", "Sessions serveur (rotation, détection de réutilisation)"],
    ["assessments / findings", "Campagnes et vulnérabilités découvertes"],
    ["post_ex_events", "Événements de post-exploitation (session, privesc, exfil…)"],
    ["schedules", "Définitions de tests récurrents (cron)"],
    ["log_entries / audit_events", "Journaux d'exécution et journal d'audit chaîné"],
    ["reports / chat_sessions / chat_messages", "Rapports et conversations de l'assistant"],
], widths=[2.7, 3.6], caption="Modèle de données (principales tables).")
h2("6.6 Conception de la sécurité")
p("La sécurité repose sur une authentification par cookies httpOnly (jeton d'accès non lisible en "
  "JavaScript), une protection CSRF par double soumission, un second facteur TOTP obligatoire "
  "pour les rôles à privilèges, une rotation des jetons de rafraîchissement avec détection de "
  "réutilisation, ainsi qu'un contrôle d'accès par rôles appliqué à chaque endpoint.")
figure("rbac.png", "Matrice RBAC : permissions accordées par rôle système.")
h2("6.7 Conclusion")
p("La conception établit une base sûre et modulaire. Le chapitre suivant décrit les technologies "
  "retenues et la réalisation.")

h1("Chapitre 7 — Technologies et réalisation")
h2("7.1 Pile technologique")
table(["Couche", "Technologies"], [
    ["Frontend", "Next.js 16 (App Router), React 19, TypeScript, Tailwind CSS, Recharts"],
    ["Backend", "FastAPI (async), SQLAlchemy 2.0, Pydantic v2, Uvicorn"],
    ["Sécurité", "Argon2, PyJWT, pyotp (TOTP), slowapi (rate-limit)"],
    ["Données", "PostgreSQL (asyncpg) / SQLite (aiosqlite), Alembic, Redis"],
    ["Planification", "APScheduler, croniter"],
    ["Moteur", "Strix (CLI multi-agent), Docker, Kali Linux, Metasploit, RAG"],
    ["Qualité", "pytest, ruff, GitHub Actions (CI)"],
], widths=[1.4, 4.9], caption="Pile technologique du projet.")
h2("7.2 Réalisation du backend")
bullets([
    ("Authentification & RBAC — ", "connexion par cookies, MFA TOTP, sessions serveur avec "
     "rotation et détection de réutilisation, permissions resource:action vérifiées par dépendance."),
    ("Pipeline d'évaluation — ", "création, lancement, annulation ; exécution en tâche de fond et "
     "diffusion d'événements d'état, de journaux et de vulnérabilités."),
    ("Temps réel — ", "bus d'événements interne pontant le runner vers les abonnés WebSocket ; "
     "réconciliation au démarrage des exécutions orphelines."),
    ("Triage & comparaison — ", "mise à jour du statut, de la gravité et de l'assignation ; "
     "comparaison des vulnérabilités entre deux scans d'une même cible."),
    ("Planification — ", "APScheduler déclenche des tests récurrents selon une expression cron."),
    ("Observabilité — ", "middleware d'identifiant de requête, sondes /readyz et /livez ; journal "
     "d'audit inaltérable chaîné par hachage."),
])
figure("pipeline.png", "Pipeline d'évaluation en 5 phases (dont la post-exploitation).")
h2("7.3 Réalisation du frontend")
p("Le frontend, en Next.js (App Router), propose des modules dédiés : tableau de bord, campagnes "
  "et vue de détail (pipeline live, console, post-exploitation, triage, comparaison), rapports, "
  "journaux, audit, utilisateurs, rôles, planification, assistant IA et paramètres. Une palette "
  "de commandes offre une navigation rapide filtrée par permissions. L'interface est responsive, "
  "thématisable (clair/sombre) et cohérente grâce à un système de composants réutilisables.")
h2("7.4 Conclusion")
p("La réalisation couvre l'ensemble des besoins, du moteur enrichi à la console d'opérations. Le "
  "chapitre suivant présente l'interface réalisée.")

h1("Chapitre 8 — Présentation de l'interface réalisée")
h2("8.1 Introduction")
p("Ce chapitre illustre, par des captures d'écran réelles, les principales vues de la console "
  "Strix Console telle qu'elle a été réalisée et vérifiée.")
h2("8.2 Authentification et tableau de bord")
figure("login.png", "Page de connexion sécurisée (cookies httpOnly, étape MFA).", base=SHOTS, width=5.7)
figure("dashboard.png", "Tableau de bord : statistiques, activité et répartition des vulnérabilités.",
       base=SHOTS)
h2("8.3 Gestion et détail des campagnes")
figure("assessments.png", "Liste des campagnes de tests d'intrusion.", base=SHOTS)
figure("assessment_detail.png",
       "Détail d'une campagne : pipeline en 5 phases, post-exploitation et triage des vulnérabilités.",
       base=SHOTS)
h2("8.4 Sortie du moteur en temps réel")
figure("engine_output.png",
       "Écran de sortie du moteur : arbre d'agents, télémétrie (jetons, coût) et console live.",
       base=SHOTS)
h2("8.5 Planification et gouvernance")
figure("schedules.png", "Planification de tests récurrents (expression cron, prochaine exécution).",
       base=SHOTS)
figure("roles.png", "Gestion des rôles et des permissions (RBAC).", base=SHOTS)
h2("8.6 Conclusion")
p("L'interface concrétise l'ensemble des besoins fonctionnels dans une expérience cohérente, "
  "sécurisée et professionnelle. Le chapitre suivant en évalue la qualité.")

# ══════════════════════════════════════════════════════════════════════════════
# PARTIE VI — VALIDATION
# ══════════════════════════════════════════════════════════════════════════════
part("Partie VI", "Tests, évaluation et perspectives")

h1("Chapitre 9 — Tests et évaluation")
h2("9.1 Environnement de test")
p("Les essais ont été menés sur un poste Windows 11 avec Docker Desktop, l'intégration ayant été "
  "validée sur une image de bac à sable Kali contenant Metasploit, contre une cible vulnérable de "
  "type application Drupal. Le backend a été exécuté en mode asynchrone (SQLite en développement, "
  "PostgreSQL en cible), le frontend via le serveur de développement Next.js.")
h2("9.2 Validation fonctionnelle")
p("Chaque incrément a fait l'objet d'une vérification de bout en bout : connexion par cookies et "
  "MFA, application effective du RBAC (retour 403 pour un observateur sur une action "
  "d'administration), exécution et suivi temps réel d'une campagne, visualisation de la sortie du "
  "moteur avec télémétrie réelle (par exemple 6 requêtes LLM et 68 542 jetons sur une exécution), "
  "triage d'une vulnérabilité, comparaison de scans, planification récurrente et export de "
  "rapports. Côté intégration moteur, exploit_research a produit, sur une cible Drupal, une "
  "recommandation ancrée (CVE-2018-7600, Drupalgeddon2) avec sélection du module Metasploit.")
h2("9.3 Couverture de tests")
p("La plateforme est couverte par un harnais de tests automatisés (pytest) exécuté en intégration "
  "continue, complété par des vérifications navigateur. L'ensemble des tests est au vert.")
figure("tests.png", "Couverture de tests par module (39 tests, 100 % réussis).")
h2("9.4 Résultats et analyse comparative")
p("À titre illustratif, la répartition des vulnérabilités par gravité sur une campagne de "
  "démonstration est présentée ci-dessous, suivie d'une analyse comparative qualitative entre un "
  "pentest manuel et l'usage de la plateforme.")
figure("severity.png", "Répartition des vulnérabilités par sévérité (campagne de démonstration).")
figure("comparison.png", "Analyse comparative : approche manuelle vs plateforme.")
table(["Critère", "Résultat"], [
    ["Tests automatisés", "39 tests, 100 % réussis"],
    ["Analyse statique (ruff)", "Aucune erreur"],
    ["Build frontend", "14+ routes compilées, 0 vulnérabilité npm"],
    ["Sécurité", "Cookies httpOnly, CSRF, MFA, RBAC, audit chaîné vérifiés"],
    ["Temps réel", "WebSocket : état, journaux et vulnérabilités diffusés"],
    ["Intégration moteur", "exploit_research validé (RAG + module Metasploit)"],
], widths=[2.2, 4.1], caption="Synthèse des résultats de validation.")
h2("9.5 Limites et perspectives")
bullets([
    ("Ordonnanceur mono-processus — ", "APScheduler s'exécute en processus unique ; un "
     "déploiement multi-workers nécessiterait un processus d'ordonnancement dédié."),
    ("Runner en thread — ", "l'exécution des scans s'appuie sur un thread interne ; une extraction "
     "vers un worker dédié améliorerait la durabilité."),
    ("Couverture du corpus RAG — ", "la qualité des réponses dépend de la fraîcheur et de "
     "l'étendue des connaissances ingérées."),
    ("Perspectives — ", "pièces jointes de preuve, scans multi-cibles, intégrations (Jira, SIEM), "
     "cartographie CVSS/CWE/OWASP, client TypeScript généré depuis l'OpenAPI."),
])
h2("9.6 Conclusion")
p("L'évaluation confirme la robustesse fonctionnelle et la qualité de la plateforme, ainsi que la "
  "valeur de l'intégration Metasploit + RAG. Les limites identifiées ouvrent des pistes "
  "d'amélioration claires.")

# ══════════════════════════════════════════════════════════════════════════════
# CONCLUSION GÉNÉRALE
# ══════════════════════════════════════════════════════════════════════════════
h1("Conclusion générale")
p("Ce projet de fin d'études a abordé l'automatisation avancée des tests d'intrusion par un "
  "système d'IA agentique fondé sur des LLM, en construisant une plateforme complète et cohérente "
  "à partir de quatre briques complémentaires. Le moteur agentique Strix fournit l'orchestration "
  "multi-agent et l'exécution des outils offensifs dans un bac à sable isolé. Le service Pentest "
  "RAG ancre le raisonnement des agents dans un corpus offensif vérifié, réduisant les "
  "hallucinations. L'intégration de Metasploit, via l'outil déterministe exploit_research, "
  "fiabilise la phase la plus délicate — l'exploitation — en déplaçant la décision du modèle vers "
  "du code vérifiable et en privilégiant un exécuteur éprouvé, sous le contrôle d'une barrière de "
  "validation humaine. Enfin, la console Strix Console transforme cet outillage en un véritable "
  "service d'opérations de sécurité, doté d'une authentification forte, d'un contrôle d'accès par "
  "rôles, d'un suivi temps réel, de la visualisation de la sortie du moteur, de la "
  "post-exploitation, du triage, de la comparaison de scans, de la planification récurrente, du "
  "reporting et d'un journal d'audit inaltérable.")
p("Développé avec Next.js et FastAPI, l'ensemble a été validé par 39 tests automatisés et vérifié "
  "de bout en bout, jusqu'à la production d'une recommandation d'exploitation ancrée "
  "(Drupalgeddon2) sur une cible réelle. Au-delà des résultats obtenus, ce travail confirme la "
  "pertinence d'une approche hybride associant l'autonomie des agents IA, la robustesse d'outils "
  "éprouvés et la rigueur d'une base de connaissances vérifiée, le tout encadré par une "
  "plateforme sûre et gouvernée. Les perspectives — extraction du runner vers une file durable, "
  "enrichissement continu du corpus RAG, intégrations métier (Jira, SIEM), cartographie normative "
  "des vulnérabilités et génération automatique du client d'API — dessinent une évolution "
  "naturelle vers un produit d'entreprise.")

# ── Bibliographie ─────────────────────────────────────────────────────────────
h1("Bibliographie & Webographie")
refs = [
    "Strix — Moteur multi-agent open-source de tests d'intrusion. https://github.com/usestrix/strix",
    "Rapid7 — Metasploit Framework. https://www.metasploit.com",
    "Qdrant — Base de données vectorielle. https://qdrant.tech",
    "LangChain — Framework d'orchestration LLM. https://python.langchain.com",
    "Sentence-Transformers / Cross-Encoders. https://www.sbert.net",
    "Nomic Embed — nomic-ai/nomic-embed-text-v1. https://huggingface.co/nomic-ai",
    "OWASP — Web Security Testing Guide et OWASP Top 10. https://owasp.org",
    "PentestGPT — LLM-driven penetration testing. Travaux de recherche associés.",
    "FastAPI — Framework web asynchrone Python. https://fastapi.tiangolo.com",
    "Next.js — Framework React (App Router). https://nextjs.org",
    "SQLAlchemy 2.0 — ORM Python. https://www.sqlalchemy.org",
    "APScheduler — Ordonnancement de tâches en Python. https://apscheduler.readthedocs.io",
    "MITRE — Common Vulnerabilities and Exposures (CVE). https://cve.mitre.org",
    "Retrieval-Augmented Generation (RAG) — Lewis et al., 2020.",
]
for i, r in enumerate(refs, 1):
    para = doc.add_paragraph(); para.paragraph_format.space_after = Pt(4)
    para.add_run(f"[{i}] ").bold = True; para.add_run(r)

out = os.path.join(HERE, "Rapport_PFE_Plateforme_Strix_Complet.docx")
doc.save(out)
print("SAVED:", out)
print("Figures:", _fig[0], "| Tableaux:", _tbl[0], "| Paragraphes:", len(doc.paragraphs))
